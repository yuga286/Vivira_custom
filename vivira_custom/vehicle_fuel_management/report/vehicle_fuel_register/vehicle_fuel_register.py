import calendar
import json
from collections import defaultdict
from datetime import date
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import cint, escape_html, flt, formatdate, getdate


MONTHS = list(calendar.month_name)[1:]
DEFAULT_FUEL_SECTIONS = ("DIESEL", "PETROL")


def execute(filters=None):
	filters = normalize_filters(filters)
	context = get_report_context(filters)
	return get_columns(), [{"html": render_report_html(context)}]


def get_columns():
	return [{"label": _("Vehicle Fuel Register"), "fieldname": "html", "fieldtype": "HTML", "width": 300}]


def normalize_filters(filters=None):
	filters = frappe._dict(filters or {})
	if not filters.get("company"):
		frappe.throw(_("Company is required."))
	if not filters.get("project"):
		frappe.throw(_("Project is required."))

	month = filters.get("month") or MONTHS[getdate().month - 1]
	if isinstance(month, int) or str(month).isdigit():
		month_index = cint(month)
		if month_index < 1 or month_index > 12:
			frappe.throw(_("Month must be between 1 and 12."))
		month = MONTHS[month_index - 1]
	else:
		try:
			month_index = MONTHS.index(str(month).title()) + 1
			month = MONTHS[month_index - 1]
		except ValueError:
			frappe.throw(_("Invalid month {0}.").format(month))

	year = cint(filters.get("year") or getdate().year)
	_, last_day = calendar.monthrange(year, month_index)
	filters.month = month
	filters.month_index = month_index
	filters.year = year
	filters.from_date = date(year, month_index, 1)
	filters.to_date = date(year, month_index, last_day)
	filters.days = [date(year, month_index, day) for day in range(1, last_day + 1)]
	return filters


def get_report_context(filters):
	sections = get_sections(filters)
	return frappe._dict(
		company_heading=get_company_heading(filters),
		month_label=f"{filters.month}-{str(filters.year)[-2:]}",
		filters=filters,
		days=filters.days,
		sections=[build_fuel_context(filters, section) for section in sections],
	)


def get_company_heading(filters):
	company = filters.get("company") or frappe.db.get_value("Project", filters.project, "company") or ""
	project_name = frappe.db.get_value("Project", filters.project, "project_name") or filters.project
	return " ".join(value for value in [company, project_name] if value)


def build_fuel_context(filters, section):
	opening_by_day, received_by_day, issue_by_day, balance_by_day = get_daily_balance(filters, section.item)
	vehicles, vehicle_totals = get_vehicle_matrix(filters, section)
	total_received = sum(flt(received_by_day.get(day)) for day in filters.days)
	total_issued = sum(flt(vehicle_totals.get(day)) for day in filters.days)
	opening_stock = flt(opening_by_day.get(filters.days[0])) if filters.days else 0
	closing_stock = opening_stock + total_received - total_issued

	return frappe._dict(
		label=section.label,
		item=section.item,
		register_title=_("{0} REGISTER (FOR THE MONTH OF {1})").format(section.label, f"{filters.month}-{str(filters.year)[-2:]}"),
		vehicle_heading=_("Used By") if section.label == "PETROL" else _("Name of Vehicle"),
		opening_by_day=opening_by_day,
		received_by_day=received_by_day,
		issue_by_day=vehicle_totals,
		balance_by_day=balance_by_day,
		vehicles=vehicles,
		totals_by_day=vehicle_totals,
		total_received=total_received,
		total_issued=total_issued,
		opening_stock=opening_stock,
		closing_stock=closing_stock,
		received_rows=get_received_rows(filters, section),
	)


def get_daily_balance(filters, fuel_type):
	opening_balance = get_opening_balance(filters, fuel_type)
	received_by_day = get_daily_received(filters, fuel_type)
	issue_by_day = get_daily_issued(filters, fuel_type)
	opening_by_day = {}
	balance_by_day = {}
	running_balance = opening_balance

	for day in filters.days:
		opening_by_day[day] = running_balance
		running_balance = running_balance + flt(received_by_day.get(day)) - flt(issue_by_day.get(day))
		balance_by_day[day] = running_balance

	return opening_by_day, received_by_day, issue_by_day, balance_by_day


def get_opening_balance(filters, fuel_type):
	company_condition, values = get_company_condition(filters)
	vehicle_condition, vehicle_values = get_vehicle_condition(filters)
	values.update(vehicle_values)
	values.update({"project": filters.project, "fuel_type": fuel_type, "from_date": filters.from_date})
	received = frappe.db.sql(
		f"""
		select coalesce(sum(qty), 0)
		from `tabFuel Receipt`
		where docstatus = 1
			and project = %(project)s
			and fuel_type = %(fuel_type)s
			and posting_date < %(from_date)s
			{company_condition}
		""",
		values,
	)[0][0]
	issued = frappe.db.sql(
		f"""
		select coalesce(sum(qty_issued), 0)
		from `tabFuel Issue`
		where docstatus = 1
			and project = %(project)s
			and fuel_type = %(fuel_type)s
			and posting_date < %(from_date)s
			{company_condition}
			{vehicle_condition}
		""",
		values,
	)[0][0]
	return flt(received) - flt(issued)


def get_daily_received(filters, fuel_type):
	company_condition, values = get_company_condition(filters)
	values.update(
		{
			"project": filters.project,
			"fuel_type": fuel_type,
			"from_date": filters.from_date,
			"to_date": filters.to_date,
		}
	)
	rows = frappe.db.sql(
		f"""
		select posting_date, coalesce(sum(qty), 0) as qty
		from `tabFuel Receipt`
		where docstatus = 1
			and project = %(project)s
			and fuel_type = %(fuel_type)s
			and posting_date between %(from_date)s and %(to_date)s
			{company_condition}
		group by posting_date
		""",
		values,
		as_dict=True,
	)
	return {getdate(row.posting_date): flt(row.qty) for row in rows}


def get_daily_issued(filters, fuel_type):
	company_condition, values = get_company_condition(filters)
	vehicle_condition, vehicle_values = get_vehicle_condition(filters)
	values.update(vehicle_values)
	values.update(
		{
			"project": filters.project,
			"fuel_type": fuel_type,
			"from_date": filters.from_date,
			"to_date": filters.to_date,
		}
	)
	rows = frappe.db.sql(
		f"""
		select posting_date, coalesce(sum(qty_issued), 0) as qty
		from `tabFuel Issue`
		where docstatus = 1
			and project = %(project)s
			and fuel_type = %(fuel_type)s
			and posting_date between %(from_date)s and %(to_date)s
			{company_condition}
			{vehicle_condition}
		group by posting_date
		""",
		values,
		as_dict=True,
	)
	return {getdate(row.posting_date): flt(row.qty) for row in rows}


def get_vehicle_matrix(filters, section):
	company_condition, values = get_company_condition(filters)
	vehicle_condition, vehicle_values = get_vehicle_condition(filters)
	values.update(vehicle_values)
	values.update(
		{
			"project": filters.project,
			"fuel_type": section.item,
			"from_date": filters.from_date,
			"to_date": filters.to_date,
		}
	)
	rows = frappe.db.sql(
		f"""
		select
			name, posting_date, fuel_type, vehicle, vehicle_no, vehicle_name, vehicle_status,
			owner_name, supplier, employee, employee_name, qty_issued
		from `tabFuel Issue`
		where docstatus = 1
			and project = %(project)s
			and fuel_type = %(fuel_type)s
			and posting_date between %(from_date)s and %(to_date)s
			{company_condition}
			{vehicle_condition}
		order by vehicle_no asc, vehicle asc, posting_date asc, posting_time asc, creation asc
		""",
		values,
		as_dict=True,
	)

	vehicles = {}
	issue_totals = defaultdict(float)
	for row in rows:
		key = (row.vehicle or row.vehicle_no or row.name, row.fuel_type)
		vehicle = vehicles.setdefault(
			key,
			frappe._dict(
				owner_name=row.owner_name or row.supplier,
				status=row.vehicle_status,
				vehicle_no=row.vehicle_no or row.vehicle,
				vehicle_or_used_by=get_vehicle_display_value(row, section.label),
				start_date=row.posting_date,
				issued_type=row.fuel_type,
				qty_by_day=defaultdict(float),
			),
		)
		vehicle.start_date = min(getdate(vehicle.start_date), getdate(row.posting_date))
		vehicle.qty_by_day[getdate(row.posting_date)] += flt(row.qty_issued)
		issue_totals[getdate(row.posting_date)] += flt(row.qty_issued)

	data = []
	for idx, vehicle in enumerate(vehicles.values(), start=1):
		total = sum(flt(vehicle.qty_by_day.get(day)) for day in filters.days)
		data.append(
			frappe._dict(
				sl_no=idx,
				owner_name=vehicle.owner_name,
				status=vehicle.status,
				vehicle_no=vehicle.vehicle_no,
				vehicle_or_used_by=vehicle.vehicle_or_used_by,
				start_date=vehicle.start_date,
				issued_type=vehicle.issued_type,
				qty_by_day=vehicle.qty_by_day,
				total=total,
			)
		)

	return data, issue_totals


def get_vehicle_display_value(row, fuel_label):
	if fuel_label == "PETROL":
		return row.employee_name or row.employee or row.vehicle_name or row.vehicle
	return row.vehicle_name or row.employee_name or row.employee or row.vehicle


def get_received_rows(filters, section):
	company_condition, values = get_company_condition(filters)
	values.update(
		{
			"project": filters.project,
			"fuel_type": section.item,
			"from_date": filters.from_date,
			"to_date": filters.to_date,
		}
	)
	rows = frappe.db.sql(
		f"""
		select posting_date, supplier, supplier_name, qty, unit_rate, total_amount
		from `tabFuel Receipt`
		where docstatus = 1
			and project = %(project)s
			and fuel_type = %(fuel_type)s
			and posting_date between %(from_date)s and %(to_date)s
			{company_condition}
		order by posting_date asc, posting_time asc, creation asc
		""",
		values,
		as_dict=True,
	)
	for row in rows:
		row.amount = flt(row.total_amount) or flt(row.qty) * flt(row.unit_rate)
	return rows


def get_sections(filters):
	if filters.get("fuel_type"):
		return [frappe._dict(label=get_fuel_label(filters.fuel_type), item=filters.fuel_type)]

	return [
		frappe._dict(label=label, item=find_fuel_item(label, filters) or label)
		for label in DEFAULT_FUEL_SECTIONS
	]


def find_fuel_item(label, filters):
	item = frappe.db.get_value("Item", {"name": label}, "name")
	if item:
		return item

	item = frappe.db.sql(
		"""
		select name
		from `tabItem`
		where disabled = 0 and (name like %(label)s or item_name like %(label)s)
		order by name
		limit 1
		""",
		{"label": f"%{label}%"},
	)
	if item:
		return item[0][0]

	values = {"project": filters.project, "label": f"%{label}%"}
	company_condition = "and company = %(company)s" if filters.get("company") else ""
	if filters.get("company"):
		values["company"] = filters.company
	for doctype in ("Fuel Receipt", "Fuel Issue"):
		row = frappe.db.sql(
			f"""
			select fuel_type
			from `tab{doctype}`
			where docstatus = 1
				and project = %(project)s
				and fuel_type like %(label)s
				{company_condition}
			limit 1
			""",
			values,
		)
		if row:
			return row[0][0]
	return None


def get_fuel_label(fuel_type):
	item_name = frappe.db.get_value("Item", fuel_type, "item_name") or fuel_type
	label = item_name.upper()
	if "DIESEL" in label:
		return "DIESEL"
	if "PETROL" in label:
		return "PETROL"
	return label


def render_report_html(context):
	parts = [get_report_style(), '<div class="vehicle-fuel-register-wrapper"><table class="vehicle-fuel-register">']
	diesel = next((section for section in context.sections if section.label == "DIESEL"), None)

	if diesel:
		parts.append(render_diesel_section(context, diesel))
	parts.append("</table></div>")
	return "".join(parts)


def get_report_style():
	return """
	<style>
		.vehicle-fuel-register-wrapper { overflow-x: auto; padding: 0; }
		.vehicle-fuel-register { border-collapse: collapse; width: max-content; min-width: 100%; font-size: 11px; color: var(--text-color); }
		.vehicle-fuel-register th, .vehicle-fuel-register td { border: 1px solid #222; padding: 3px 5px; line-height: 1.2; height: 23px; white-space: nowrap; vertical-align: middle; }
		.vehicle-fuel-register th { font-weight: 700; text-align: center; background: #fff; }
		.vehicle-fuel-register .main-heading, .vehicle-fuel-register .section-heading { font-weight: 700; text-align: center; font-size: 12px; }
		.vehicle-fuel-register .metric-label, .vehicle-fuel-register .total-row td, .vehicle-fuel-register .summary-title { font-weight: 700; }
		.vehicle-fuel-register .num { text-align: center; }
		.vehicle-fuel-register .text { text-align: left; }
		.vehicle-fuel-register .summary-label { font-weight: 600; min-width: 145px; }
		.vehicle-fuel-register .summary-value { min-width: 80px; text-align: right; }
		.vehicle-fuel-register .section-gap td { border: 0; height: 20px; }
		.vehicle-fuel-register .sticky-col { position: sticky; background: #fff; z-index: 2; }
		.vehicle-fuel-register .col-sl { left: 0; min-width: 45px; }
		.vehicle-fuel-register .col-owner { left: 45px; min-width: 165px; }
		.vehicle-fuel-register .col-status { left: 210px; min-width: 90px; }
		.vehicle-fuel-register .col-vehicle-no { left: 300px; min-width: 115px; }
		.vehicle-fuel-register .col-vehicle { left: 415px; min-width: 145px; }
		.vehicle-fuel-register .col-start { left: 560px; min-width: 90px; }
		.vehicle-fuel-register .col-issued-type { left: 650px; min-width: 115px; }
		.vehicle-fuel-register-report .dt-scrollable { display: none; }
		.vehicle-fuel-register-report .report-html { overflow: visible; }
		@media print {
			@page { size: landscape; margin: 8mm; }
			.vehicle-fuel-register-wrapper { overflow: visible; }
			.vehicle-fuel-register { font-size: 8px; }
			.vehicle-fuel-register th, .vehicle-fuel-register td { padding: 2px 3px; height: 18px; }
			.vehicle-fuel-register .sticky-col { position: static; }
		}
	</style>
	"""


def render_diesel_section(context, section):
	days = context.days
	summary = [
		("Opeing stock", section.opening_stock),
		("Total Qty Received", section.total_received),
		("Total Qty Issued", section.total_issued),
		("Closing Stock", section.closing_stock),
	]
	rows = [
		"<tr>"
		+ cell(context.company_heading, "td", 'class="main-heading" colspan="7"')
		+ cell("Date", "td", 'class="metric-label num"')
		+ "".join(cell(formatdate(day), "td", 'class="num"') for day in days)
		+ cell("", "td")
		+ cell("", "td")
		+ cell("", "td")
		+ "</tr>",
		"<tr>"
		+ cell(section.register_title, "td", 'class="section-heading" colspan="7"')
		+ cell("OPENING BALANCE", "td", 'class="metric-label num"')
		+ render_value_cells(days, section.opening_by_day, zero_as_blank=False)
		+ cell("TOTAL", "td", 'class="summary-title num"')
		+ cell(summary[0][0], "td", 'class="summary-label"')
		+ cell(fmt_number(summary[0][1]), "td", 'class="summary-value"')
		+ "</tr>",
		render_header_and_metric_row(context, section, "Received", section.received_by_day, summary[1]),
		render_blank_metric_row(days, "Issue", section.issue_by_day, summary[2]),
		render_blank_metric_row(days, "BALANCE", section.balance_by_day, summary[3]),
	]
	rows.extend(render_vehicle_rows(context, section))
	rows.append(render_total_row(context, section))
	return "".join(rows)


def render_vehicle_section(context, section):
	rows = [
		"<tr>"
		+ cell(section.register_title, "td", 'class="section-heading" colspan="7"')
		+ cell("", "td")
		+ "".join(cell("", "td") for _day in context.days)
		+ cell("", "td")
		+ "</tr>",
		render_vehicle_header_row(context, section),
	]
	rows.extend(render_vehicle_rows(context, section))
	rows.append(render_total_row(context, section))
	return "".join(rows)


def render_header_and_metric_row(context, section, label, values_by_day, summary_item):
	return (
		"<tr>"
		+ render_identity_header_cells(section)
		+ cell(label, "td", 'class="metric-label num"')
		+ render_value_cells(context.days, values_by_day, zero_as_blank=False)
		+ cell("", "td")
		+ cell(summary_item[0], "td", 'class="summary-label"')
		+ cell(fmt_number(summary_item[1]), "td", 'class="summary-value"')
		+ "</tr>"
	)


def render_blank_metric_row(days, label, values_by_day, summary_item):
	return (
		"<tr>"
		+ "".join(cell("", "td") for _idx in range(7))
		+ cell(label, "td", 'class="metric-label num"')
		+ render_value_cells(days, values_by_day, zero_as_blank=False)
		+ cell("", "td")
		+ cell(summary_item[0], "td", 'class="summary-label"')
		+ cell(fmt_number(summary_item[1]), "td", 'class="summary-value"')
		+ "</tr>"
	)


def render_vehicle_header_row(context, section):
	return (
		"<tr>"
		+ render_identity_header_cells(section)
		+ cell("Date", "th")
		+ "".join(cell(day.strftime("%d"), "th", 'class="num"') for day in context.days)
		+ cell("TOTAL", "th")
		+ "</tr>"
	)


def render_identity_header_cells(section):
	return "".join(
		[
			cell("Sl.No", "th", 'class="sticky-col col-sl"'),
			cell("Name of Owner", "th", 'class="sticky-col col-owner"'),
			cell("Status", "th", 'class="sticky-col col-status"'),
			cell("Vehicle No", "th", 'class="sticky-col col-vehicle-no"'),
			cell(section.vehicle_heading, "th", 'class="sticky-col col-vehicle"'),
			cell("Start Date", "th", 'class="sticky-col col-start"'),
			cell("Type of Issued", "th", 'class="sticky-col col-issued-type"'),
		]
	)


def render_vehicle_rows(context, section):
	rows = []
	for vehicle in section.vehicles:
		rows.append(
			"<tr>"
			+ cell(vehicle.sl_no, "td", 'class="num sticky-col col-sl"')
			+ cell(vehicle.owner_name, "td", 'class="text sticky-col col-owner"')
			+ cell(vehicle.status, "td", 'class="text sticky-col col-status"')
			+ cell(vehicle.vehicle_no, "td", 'class="text sticky-col col-vehicle-no"')
			+ cell(vehicle.vehicle_or_used_by, "td", 'class="text sticky-col col-vehicle"')
			+ cell(formatdate(vehicle.start_date) if vehicle.start_date else "", "td", 'class="num sticky-col col-start"')
			+ cell(vehicle.issued_type, "td", 'class="text sticky-col col-issued-type"')
			+ cell("", "td")
			+ render_value_cells(context.days, vehicle.qty_by_day, zero_as_blank=True)
			+ cell(fmt_number(vehicle.total), "td", 'class="num"')
			+ "</tr>"
		)
	return rows


def render_total_row(context, section):
	return (
		'<tr class="total-row">'
		+ cell("Total", "td", 'colspan="8"')
		+ render_value_cells(context.days, section.totals_by_day, zero_as_blank=False)
		+ cell(fmt_number(section.total_issued), "td", 'class="num"')
		+ "</tr>"
	)


def render_received_section(context, section):
	total_qty = sum(flt(row.qty) for row in section.received_rows)
	total_amount = sum(flt(row.amount) for row in section.received_rows)
	rows = [
		"<tr>"
		+ cell("", "td")
		+ cell(_("{0} INWARD/RECEIVED (FOR THE MONTH OF {1})").format(section.label, context.month_label), "td", 'class="section-heading" colspan="5"')
		+ "".join(cell("", "td") for _idx in range(len(context.days) + 4))
		+ "</tr>",
		"<tr>"
		+ cell("", "td")
		+ cell("Date", "th")
		+ cell("Vendor/Party Name", "th")
		+ cell("Qty", "th")
		+ cell("Rate", "th")
		+ cell("Amount", "th")
		+ "".join(cell("", "td") for _idx in range(len(context.days) + 4))
		+ "</tr>",
	]
	for row in section.received_rows:
		rows.append(
			"<tr>"
			+ cell("", "td")
			+ cell(formatdate(row.posting_date), "td", 'class="num"')
			+ cell(row.supplier_name or row.supplier, "td", 'class="text"')
			+ cell(fmt_number(row.qty), "td", 'class="num"')
			+ cell(fmt_number(row.unit_rate), "td", 'class="num"')
			+ cell(fmt_number(row.amount), "td", 'class="num"')
			+ "".join(cell("", "td") for _idx in range(len(context.days) + 4))
			+ "</tr>"
		)
	rows.append(
		'<tr class="total-row">'
		+ cell("", "td")
		+ cell("Total", "td", 'colspan="2"')
		+ cell(fmt_number(total_qty), "td", 'class="num"')
		+ cell("", "td")
		+ cell(fmt_number(total_amount), "td", 'class="num"')
		+ "".join(cell("", "td") for _idx in range(len(context.days) + 4))
		+ "</tr>"
	)
	return "".join(rows)


def render_gap():
	return '<tr class="section-gap"><td colspan="50"></td></tr>'


def render_value_cells(days, values_by_day, zero_as_blank):
	cells = []
	for day in days:
		value = flt(values_by_day.get(day))
		cells.append(cell("" if zero_as_blank and not value else fmt_number(value), "td", 'class="num"'))
	return "".join(cells)


def cell(value, tag="td", attrs=""):
	attrs = f" {attrs}" if attrs else ""
	return f"<{tag}{attrs}>{escape_html(value if value is not None else '')}</{tag}>"


def fmt_number(value):
	value = flt(value)
	if value == int(value):
		return int(value)
	return round(value, 2)


def get_company_condition(filters):
	values = {}
	if filters.get("company"):
		values["company"] = filters.company
		return "and company = %(company)s", values
	return "", values


def get_vehicle_condition(filters):
	values = {}
	if filters.get("vehicle"):
		values["vehicle"] = filters.vehicle
		return "and vehicle = %(vehicle)s", values
	return "", values


@frappe.whitelist()
def download_excel(filters):
	filters = normalize_filters(json.loads(filters or "{}"))
	context = get_report_context(filters)
	workbook = make_workbook(context)
	stream = BytesIO()
	workbook.save(stream)
	stream.seek(0)

	frappe.response["filename"] = f"Vehicle Fuel Register {filters.month}-{filters.year}.xlsx"
	frappe.response["filecontent"] = stream.read()
	frappe.response["type"] = "binary"


def make_workbook(context):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter

	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "Vehicle Fuel Register"
	sheet.sheet_view.showGridLines = False
	thin = Side(style="thin", color="000000")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	bold = Font(bold=True)
	center = Alignment(horizontal="center", vertical="center")
	left = Alignment(horizontal="left", vertical="center")
	heading_fill = PatternFill("solid", fgColor="FFFFFF")
	row_idx = 1

	diesel = next((section for section in context.sections if section.label == "DIESEL"), None)
	if diesel:
		row_idx = write_diesel_sheet_section(sheet, context, diesel, row_idx, border, bold, center, left, heading_fill)

	for col_idx in range(1, len(context.days) + 12):
		width = 5 if col_idx == 1 else 10
		if col_idx == 2:
			width = 25
		elif col_idx in (3, 4, 5):
			width = 16
		elif col_idx in (6, 7):
			width = 12
		elif 9 <= col_idx <= len(context.days) + 8:
			width = 8
		sheet.column_dimensions[get_column_letter(col_idx)].width = width

	sheet.freeze_panes = "H1"
	sheet.page_setup.orientation = "landscape"
	sheet.page_setup.fitToWidth = 1
	return workbook


def write_diesel_sheet_section(sheet, context, section, row_idx, border, bold, center, left, heading_fill):
	days = context.days
	summary = [
		("Opeing stock", section.opening_stock),
		("Total Qty Received", section.total_received),
		("Total Qty Issued", section.total_issued),
		("Closing Stock", section.closing_stock),
	]
	write_merged(sheet, row_idx, 1, row_idx, 7, context.company_heading, bold, center, border, heading_fill)
	write_cell(sheet, row_idx, 8, "Date", bold, center, border)
	write_day_headers(sheet, row_idx, days, 9, border, center)
	row_idx += 1

	write_merged(sheet, row_idx, 1, row_idx, 7, section.register_title, bold, center, border, heading_fill)
	write_metric_sheet_row(sheet, row_idx, days, "OPENING BALANCE", section.opening_by_day, summary[0], border, bold, center)
	row_idx += 1
	write_identity_sheet_headers(sheet, row_idx, section, border, bold, center)
	write_metric_sheet_row(sheet, row_idx, days, "Received", section.received_by_day, summary[1], border, bold, center)
	row_idx += 1
	write_metric_sheet_row(sheet, row_idx, days, "Issue", section.issue_by_day, summary[2], border, bold, center)
	row_idx += 1
	write_metric_sheet_row(sheet, row_idx, days, "BALANCE", section.balance_by_day, summary[3], border, bold, center)
	row_idx += 1
	row_idx = write_vehicle_sheet_rows(sheet, context, section, row_idx, border, bold, center, left)
	return write_total_sheet_row(sheet, context, section, row_idx, border, bold, center)


def write_vehicle_sheet_section(sheet, context, section, row_idx, border, bold, center, left):
	write_merged(sheet, row_idx, 1, row_idx, 7, section.register_title, bold, center, border)
	row_idx += 1
	write_identity_sheet_headers(sheet, row_idx, section, border, bold, center)
	write_cell(sheet, row_idx, 8, "Date", bold, center, border)
	write_day_headers(sheet, row_idx, context.days, 9, border, center)
	write_cell(sheet, row_idx, len(context.days) + 9, "TOTAL", bold, center, border)
	row_idx += 1
	row_idx = write_vehicle_sheet_rows(sheet, context, section, row_idx, border, bold, center, left)
	return write_total_sheet_row(sheet, context, section, row_idx, border, bold, center)


def write_identity_sheet_headers(sheet, row_idx, section, border, bold, center):
	for col_idx, label in enumerate(
		("Sl.No", "Name of Owner", "Status", "Vehicle No", section.vehicle_heading, "Start Date", "Type of Issued"),
		start=1,
	):
		write_cell(sheet, row_idx, col_idx, label, bold, center, border)


def write_metric_sheet_row(sheet, row_idx, days, label, values_by_day, summary, border, bold, center):
	write_cell(sheet, row_idx, 8, label, bold, center, border)
	for col_idx, day in enumerate(days, start=9):
		write_cell(sheet, row_idx, col_idx, flt(values_by_day.get(day)), None, center, border)
	total_col = len(days) + 9
	write_cell(sheet, row_idx, total_col + 1, summary[0], bold, center, border)
	write_cell(sheet, row_idx, total_col + 2, summary[1], None, center, border)


def write_day_headers(sheet, row_idx, days, start_col, border, center):
	for col_idx, day in enumerate(days, start=start_col):
		write_cell(sheet, row_idx, col_idx, day, None, center, border)


def write_vehicle_sheet_rows(sheet, context, section, row_idx, border, bold, center, left):
	for vehicle in section.vehicles:
		values = [
			vehicle.sl_no,
			vehicle.owner_name,
			vehicle.status,
			vehicle.vehicle_no,
			vehicle.vehicle_or_used_by,
			vehicle.start_date,
			vehicle.issued_type,
		]
		for col_idx, value in enumerate(values, start=1):
			write_cell(sheet, row_idx, col_idx, value, None, left if col_idx in (2, 3, 4, 5, 7) else center, border)
		write_cell(sheet, row_idx, 8, "", None, center, border)
		for col_idx, day in enumerate(context.days, start=9):
			value = flt(vehicle.qty_by_day.get(day))
			write_cell(sheet, row_idx, col_idx, value if value else None, None, center, border)
		write_cell(sheet, row_idx, len(context.days) + 9, vehicle.total, None, center, border)
		row_idx += 1
	return row_idx


def write_total_sheet_row(sheet, context, section, row_idx, border, bold, center):
	write_merged(sheet, row_idx, 1, row_idx, 8, "Total", bold, center, border)
	for col_idx, day in enumerate(context.days, start=9):
		write_cell(sheet, row_idx, col_idx, flt(section.totals_by_day.get(day)), bold, center, border)
	write_cell(sheet, row_idx, len(context.days) + 9, section.total_issued, bold, center, border)
	return row_idx + 1


def write_received_sheet_section(sheet, context, section, row_idx, border, bold, center, left):
	write_merged(
		sheet,
		row_idx,
		2,
		row_idx,
		6,
		_("{0} INWARD/RECEIVED (FOR THE MONTH OF {1})").format(section.label, context.month_label),
		bold,
		center,
		border,
	)
	row_idx += 1
	for col_idx, label in enumerate(("Date", "Vendor/Party Name", "Qty", "Rate", "Amount"), start=2):
		write_cell(sheet, row_idx, col_idx, label, bold, center, border)
	row_idx += 1
	total_qty = 0
	total_amount = 0
	for receipt in section.received_rows:
		total_qty += flt(receipt.qty)
		total_amount += flt(receipt.amount)
		values = [receipt.posting_date, receipt.supplier_name or receipt.supplier, receipt.qty, receipt.unit_rate, receipt.amount]
		for col_idx, value in enumerate(values, start=2):
			write_cell(sheet, row_idx, col_idx, value, None, left if col_idx == 3 else center, border)
		row_idx += 1
	write_merged(sheet, row_idx, 2, row_idx, 3, "Total", bold, center, border)
	write_cell(sheet, row_idx, 4, total_qty, bold, center, border)
	write_cell(sheet, row_idx, 5, "", bold, center, border)
	write_cell(sheet, row_idx, 6, total_amount, bold, center, border)
	return row_idx + 1


def write_cell(sheet, row_idx, col_idx, value, font, alignment, border, fill=None):
	cell_obj = sheet.cell(row=row_idx, column=col_idx, value=value)
	if font:
		cell_obj.font = font
	if alignment:
		cell_obj.alignment = alignment
	if border:
		cell_obj.border = border
	if fill:
		cell_obj.fill = fill
	return cell_obj


def write_merged(sheet, start_row, start_col, end_row, end_col, value, font, alignment, border, fill=None):
	sheet.merge_cells(start_row=start_row, start_column=start_col, end_row=end_row, end_column=end_col)
	cell_obj = write_cell(sheet, start_row, start_col, value, font, alignment, border, fill)
	for row in range(start_row, end_row + 1):
		for col in range(start_col, end_col + 1):
			sheet.cell(row=row, column=col).border = border
	return cell_obj
